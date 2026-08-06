from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import ApplicationRecord


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_FONT = Font(color="FFFFFF", bold=True)
HEADER_FONT = Font(bold=True)
THIN_GRAY = Side(style="thin", color="D9E1F2")


def _write_sheet(ws, rows: list[dict[str, Any]], *, table_filter: bool = True) -> None:
    if not rows:
        ws.append(["暂无数据"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(key) for key in headers])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if table_filter:
        ws.auto_filter.ref = ws.dimensions

    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for cell in ws[get_column_letter(col_idx)]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 100))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
        text_heavy = any(token in str(header) for token in ["评语", "评价", "摘要", "概况", "内容", "依据", "原文", "警告"])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 45 if text_heavy else 30)
    ws.row_dimensions[1].height = 30


def _summary_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        row = record.to_summary_row()
        row["AI优势"] = "；".join(record.ai_review.strengths)
        row["AI不足"] = "；".join(record.ai_review.weaknesses)
        row["AI风险"] = "；".join(record.ai_review.risks)
        rows.append(row)
    return rows


def _education_work_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in records:
        for item in record.education:
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "类型": "教育经历",
                    "层次/职称": item.level,
                    "单位": item.institution,
                    "院系/专业": item.major,
                    "开始": item.start,
                    "结束": item.end,
                    "原文": item.raw,
                }
            )
        for item in record.work_history:
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "类型": "博士后" if item.category == "postdoc" else "工作经历",
                    "层次/职称": item.title,
                    "单位": item.institution,
                    "院系/专业": item.department,
                    "开始": item.start,
                    "结束": item.end,
                    "原文": item.raw,
                }
            )
    return rows


def _funding_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in records:
        for item in record.funding:
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "来源": item.source_type,
                    "资助机构": item.agency,
                    "项目类别": item.program_type,
                    "批准号": item.grant_no,
                    "项目名称": item.title,
                    "开始": item.start,
                    "结束": item.end,
                    "经费(万元)": item.amount_wan,
                    "状态": item.status,
                    "角色": item.role,
                    "原文": item.raw,
                }
            )
    return rows


def _publication_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in records:
        for item in record.publications:
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "序号": item.sequence,
                    "代表性论著": "是" if item.is_representative else "否",
                    "年份": item.year,
                    "期刊": item.journal,
                    "题目": item.title,
                    "本人署名": item.role_label,
                    "第一作者": "是" if item.is_first_author else "否",
                    "通讯作者": "是" if item.is_corresponding_author else "否",
                    "完整引文": item.citation,
                    "本人贡献": item.applicant_contribution,
                }
            )
    return rows


def _budget_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        b = record.budget
        rows.append(
            {
                "文件名": record.source_file,
                "申请人": record.applicant_name,
                "项目名称": record.project_title,
                "直接费用(万元)": b.direct_cost_wan,
                "设备费(万元)": b.equipment_wan,
                "业务费(万元)": b.business_wan,
                "劳务费(万元)": b.labor_wan,
                "合作转拨(万元)": b.transfer_wan,
                "其他来源(万元)": b.other_source_wan,
            }
        )
    return rows


def _ai_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        ai = record.ai_review
        rows.append(
            {
                "文件名": record.source_file,
                "申请人": record.applicant_name,
                "项目名称": record.project_title,
                "初评等级": ai.preliminary_grade,
                "总体评价": ai.overall_assessment,
                "优势": "\n".join(ai.strengths),
                "不足": "\n".join(ai.weaknesses),
                "风险": "\n".join(ai.risks),
                "基金主持情况评价": ai.funding_history_assessment,
                "论文发表情况评价": ai.publication_assessment,
                "证据缺口": "\n".join(ai.evidence_gaps),
                "最终评语": ai.final_review,
            }
        )
    return rows


def _manual_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        m = record.manual_review
        rows.append(
            {
                "文件名": record.source_file,
                "申请人": record.applicant_name,
                "项目名称": record.project_title,
                "科学问题": m.scientific_question_score,
                "创新性": m.innovation_score,
                "申请人与团队": m.applicant_team_score,
                "突破潜力": m.breakthrough_score,
                "预算合理性": m.budget_score,
                "总分": m.total_score,
                "综合评价": m.overall_grade,
                "资助意见": m.funding_opinion,
                "排序": m.rank,
                "评审备注": m.reviewer_notes,
                "最终评语": record.ai_review.final_review,
            }
        )
    return rows


def _evidence_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        for item in record.evidence:
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "字段": item.field_name,
                    "提取值": item.value,
                    "页码": item.page,
                    "置信度": item.confidence,
                    "方法": item.method,
                    "证据片段": item.snippet,
                }
            )
    return rows


def _section_rows(records: list[ApplicationRecord]) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        for name, text in record.sections.items():
            rows.append(
                {
                    "文件名": record.source_file,
                    "申请人": record.applicant_name,
                    "区块": name,
                    "内容": text,
                }
            )
    return rows


def build_excel(records: list[ApplicationRecord]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    sheets = [
        ("项目总表", _summary_rows(records)),
        ("教育与工作", _education_work_rows(records)),
        ("基金主持参与", _funding_rows(records)),
        ("论文发表", _publication_rows(records)),
        ("经费预算", _budget_rows(records)),
        ("AI初评", _ai_rows(records)),
        ("人工评分", _manual_rows(records)),
        ("提取证据", _evidence_rows(records)),
        ("正文区块", _section_rows(records)),
    ]

    for name, rows in sheets:
        ws = wb.create_sheet(name)
        _write_sheet(ws, rows)

    # 人工评分表提供下拉选项和离线总分公式。
    review_ws = wb["人工评分"]
    headers = {cell.value: cell.column for cell in review_ws[1]}
    grade_col = headers.get("综合评价")
    opinion_col = headers.get("资助意见")
    total_col = headers.get("总分")
    score_cols = [headers.get(x) for x in ["科学问题", "创新性", "申请人与团队", "突破潜力", "预算合理性"]]
    max_row = max(review_ws.max_row, 200)

    if grade_col:
        dv = DataValidation(type="list", formula1='"A 优,B 良,C 中,D 差"', allow_blank=True)
        review_ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(grade_col)}2:{get_column_letter(grade_col)}{max_row}")
    if opinion_col:
        dv = DataValidation(type="list", formula1='"A 优先资助,B 可资助,C 不予资助"', allow_blank=True)
        review_ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(opinion_col)}2:{get_column_letter(opinion_col)}{max_row}")
    if total_col and all(score_cols):
        total_letter = get_column_letter(total_col)
        score_letters = [get_column_letter(int(x)) for x in score_cols if x]
        for row in range(2, review_ws.max_row + 1):
            if review_ws[f"{total_letter}{row}"].value in (None, ""):
                review_ws[f"{total_letter}{row}"] = "=" + "+".join(f"{c}{row}" for c in score_letters)
        review_ws.conditional_formatting.add(
            f"{total_letter}2:{total_letter}{max_row}",
            CellIsRule(operator="greaterThanOrEqual", formula=["85"], fill=PatternFill("solid", fgColor="C6EFCE")),
        )

    # 首页说明。
    summary_ws = wb["项目总表"]
    summary_ws.sheet_view.showGridLines = False
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
